from docx import Document


def read_docx_full(file_path):
    """Read all text from a .docx file, including paragraphs, tables, and lists."""
    doc = Document(file_path)
    full_text = []

    # Read all paragraphs
    for paragraph in doc.paragraphs:
        full_text.append(paragraph.text)

    # Read all tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                full_text.append(cell.text)

    # Combine all text, stripping any extra whitespace
    return "\n".join([text.strip() for text in full_text if text.strip()])


# Replace 'your_file.docx' with the path to your .docx file
file_path = "/Users/rajamuhammedomar/latest-fyr/ForYourResearch/backend/app/pdf_parsing/naman/LSTM_Powered_Identification_of_Clickbait_Content_on_Entertainment_and_News_Websites.docx"
docx_text = read_docx_full(file_path)

# Print or send the text to ChatGPT
print(docx_text)
